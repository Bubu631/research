#!/usr/bin/env python3
"""Read immutable annual NHS source workbooks into provenance-rich CSV data."""
from __future__ import annotations
import base64
from collections import Counter
import csv
from datetime import datetime, timezone
import hashlib
import io
import json
import math
from pathlib import Path
import re
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
TRUST_GROUPS = {'Acute&Acute Community Trusts', 'Acute Specialist Trusts',
                'MH&LD, MH, LD&Community Trusts', 'Community Trusts', 'Ambulance Trusts'}
GROUP_MAP = {
    'Acute and Acute & Community Trusts': 'Acute&Acute Community Trusts',
    'Mental Health & Learning Disability and Mental Health, Learning Disability & Community Trusts':
        'MH&LD, MH, LD&Community Trusts',
    'Social Enterprises - Community': 'Social Enterprises Community',
    'Social Enterprises - Mental Health': 'Social Enterprises MH',
}
# Official year-specific technical guides, Section 8.1. This is not a guarantee
# of stable individual workers; it preserves the publisher's exclusions.
NO_HISTORY = {
    2023: {
        'RWK': 'Earlier sample-drawing errors; not comparable prior to 2023',
        'RA9': 'Earlier sample-drawing errors; not comparable prior to 2023',
        'RBN': 'Merger with RVY; new organisation in 2023',
        'RH5': 'Merger with RA4; new organisation in 2023',
        '0DE': 'Did not participate in 2022', 'NQ7': 'Did not participate in 2022',
        'QHM': 'Did not participate in 2022', 'QOC': 'Did not participate in 2022',
        'QOP': 'Did not participate in 2022', 'QUE': 'Did not participate in 2022',
    },
    2024: {
        'QOQ': 'Did not participate in 2023',
        'RX2': 'CAMHS transfer to RW1',
        'RW1': 'CAMHS transfer from RX2 and community/mental-health/LD services from R1F',
        'RY9': 'Hounslow community services transferred to RKL',
    },
}
PRIMARY_CAPTURES = {2023: '20240705035334', 2024: '20250427122552'}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def number(value):
    if value is None or value in {'', '-', '*', '<10', '<11', 'NA', 'N/A'}:
        return None
    if not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f'Unexpected nonnumeric score/count: {value!r}')
    return value


def write_csv(filename, rows):
    with (HERE/filename).open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def parse_benchmarks():
    manifest = json.loads((HERE/'archived_benchmark_manifest.json').read_text())
    records, audits = [], []
    for vintage, capture in PRIMARY_CAPTURES.items():
        source = next(r for r in manifest if r['vintage'] == vintage and r['archive_capture'] == capture)
        data = (HERE/source['local_file']).read_bytes()
        assert sha(data) == source['sha256']
        digest = base64.b32encode(hashlib.sha1(data).digest()).decode()
        assert digest == source['cdx_digest_sha1_base32'], 'Archive payload does not match CDX capture digest'
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        notes = '\n'.join(' | '.join(str(v) for v in row if v is not None)
                          for row in wb['Notes'].iter_rows(values_only=True))
        (HERE/f'benchmark_notes_{vintage}.txt').write_text(notes+'\n')
        count_before = len(records)
        for ws in wb:
            iterator = ws.iter_rows(values_only=True)
            header = next(iterator)
            if 'org_id' not in header:
                continue
            columns = {v: i for i, v in enumerate(header) if isinstance(v, str)}
            years = sorted(int(name.rsplit('_', 1)[1]) for name in columns
                           if re.fullmatch(r'PP4_2_20\d{2}', name))
            assert years == list(range(2021, vintage+1))
            for row_number, values in enumerate(iterator, 2):
                code = values[columns['org_id']]
                if code is None:
                    assert not any(v is not None for v in values)
                    continue
                code = str(code).strip()
                # Organisation grouping is the workbook sheet, as in the 2025
                # panel. Some social enterprises compare against a trust group;
                # that comparison field must not reclassify them as NHS trusts.
                group_source = ws.title
                comparison_group = str(values[columns['org_type_reporting_name']])
                group = GROUP_MAP.get(group_source, group_source)
                available = [y for y in years if number(values[columns[f'PP4_2_{y}']]) is not None]
                complete = available == years and code not in NO_HISTORY[vintage]
                for year in years:
                    value = number(values[columns[f'PP4_2_{year}']])
                    n = number(values[columns[f'PP4_2_n_{year}']])
                    assert value is None or 0 <= value <= 10
                    assert n is None or (n >= 0 and int(n) == n)
                    comparable = value is not None and code not in NO_HISTORY[vintage]
                    cells = {name: f'{get_column_letter(columns[key]+1)}{row_number}'
                             for name, key in [('burnout_score', f'PP4_2_{year}'),
                                               ('burnout_n', f'PP4_2_n_{year}')]}
                    records.append(dict(release_vintage=vintage, year=year, org_code=code,
                        org_name=values[columns['org_name']], benchmark_group=group,
                        benchmark_group_original=group_source, benchmark_group_vintage=vintage,
                        reporting_comparison_group=comparison_group,
                        burnout_score=value, burnout_n=n, is_trust=int(group in TRUST_GROUPS),
                        historically_comparable=int(comparable), result_available=int(value is not None),
                        complete_history_in_vintage=int(complete),
                        comparability_reason=NO_HISTORY[vintage].get(code,
                            'Publisher supplies this historical result' if value is not None else 'Historical result unavailable'),
                        archive_capture_utc=datetime.strptime(capture, '%Y%m%d%H%M%S').replace(tzinfo=timezone.utc).isoformat(),
                        official_source_url=source['original_url'], archive_source_url=source['url'],
                        accessed_at_utc=source['accessed_at_utc'], source_file=source['local_file'],
                        source_sha256=source['sha256'], source_sheet=ws.title, source_row=row_number,
                        source_cells=json.dumps(cells, sort_keys=True),
                        weighting=f'Historical scores use {vintage} reporting-frame occupational weights where applicable',
                        count_definition='Nominal valid PP4_2 composite respondents; not effective sample size'))
        own = records[count_before:]
        audits.append(dict(vintage=vintage, archive_capture=capture, archive_cdx_digest_matches=True,
            source_sha256=source['sha256'], workbook_created=str(wb.properties.created),
            workbook_modified=str(wb.properties.modified), rows=len(own),
            organisations=len({r['org_code'] for r in own}),
            trust_organisations=len({r['org_code'] for r in own if r['is_trust']}),
            complete_comparable_trust_histories=len({r['org_code'] for r in own if r['is_trust'] and r['complete_history_in_vintage']}),
            missing_scores=sum(r['burnout_score'] is None for r in own)))
    assert len({(r['release_vintage'], r['year'], r['org_code']) for r in records}) == len(records)
    records.sort(key=lambda r: (r['release_vintage'], r['org_code'], r['year']))
    write_csv('benchmark_vintage_panel.csv', records)
    return records, audits


def parse_annual():
    manifest = json.loads((HERE/'source_manifest.json').read_text())['sources']
    records = []
    for year in range(2021, 2025):
        source = manifest[f'detailed_{year}.zip']
        data = (HERE/f'detailed_{year}.zip').read_bytes()
        assert sha(data) == source['sha256']
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            member = next(n for n in z.namelist() if 'organisational results' in n)
            book = z.read(member)
        wb = openpyxl.load_workbook(io.BytesIO(book), read_only=True, data_only=True)
        ws = wb['PEOPLE PROMISE ELEMENTS']
        rows = list(ws.iter_rows(values_only=True))
        assert rows[4][0] == 'ODS code' and rows[4][28] == 'PP4_2 - score'
        assert rows[4][29] == 'Base (number of responses)'
        for row_number, row in enumerate(rows[5:], 6):
            if row[0] is None:
                assert not any(v is not None for v in row)
                continue
            group = GROUP_MAP.get(row[2], row[2])
            value, n = number(row[28]), number(row[29])
            assert value is None or 0 <= value <= 10
            assert n is None or n == int(n) and n >= 0
            records.append(dict(release_vintage=year, year=year, org_code=str(row[0]), org_name=row[1],
                benchmark_group=group, benchmark_group_original=row[2], burnout_score=value, burnout_n=n,
                is_trust=int(group in TRUST_GROUPS), weighting=row[5], source_url=source['url'],
                accessed_at_utc=source['accessed_at_utc'], source_zip_sha256=source['sha256'],
                source_member=member, source_member_sha256=sha(book), source_sheet=ws.title,
                source_row=row_number, score_cell=f'AC{row_number}', count_cell=f'AD{row_number}',
                workbook_modified=str(wb.properties.modified),
                vintage_note='Available annual archive version; not claimed to be exact first-publication-day bytes'))
    write_csv('annual_release_current_year_panel.csv', records)
    return records


def main():
    benchmark, audits = parse_benchmarks()
    annual = parse_annual()
    result = dict(status='passed', generated_at_utc=datetime.now(timezone.utc).isoformat(),
        primary_benchmark_vintages=audits, annual_rows=len(annual),
        annual_org_counts={str(y): sum(r['year'] == y for r in annual) for y in range(2021, 2025)},
        outputs={name: sha((HERE/name).read_bytes()) for name in
                 ['benchmark_vintage_panel.csv', 'annual_release_current_year_panel.csv']},
        parser_sha256=sha(Path(__file__).read_bytes()),
        scope='Archived official benchmark files provide genuine reporting vintages; availability is established by archive capture, not exact release day. No model fitting is performed.')
    (HERE/'parsing_manifest.json').write_text(json.dumps(result, indent=2)+'\n')
    print(json.dumps(result, indent=2))


if __name__ == '__main__':
    main()
