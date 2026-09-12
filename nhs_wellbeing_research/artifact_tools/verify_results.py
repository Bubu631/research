"""Read-only audit of saved XLSX values against scientific CSV sources."""
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
ART=ROOT/'artifact_tools'
OUTPUT=Path(os.environ.get('NHS_RESULTS_OUTPUT',str(ROOT.parent/'output/release')))/'nhs_wellbeing_results.xlsx'
def read(relative):return list(csv.DictReader((ROOT/relative).open()))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

manifest=json.loads((ART/'workbook_manifest.json').read_text())
assert sha(OUTPUT)==manifest['sha256']
for entry in manifest['sources']:assert sha(ROOT/entry['file'])==entry['sha256']
wb=openpyxl.load_workbook(OUTPUT,read_only=True,data_only=False)
assert wb.sheetnames==['Forecasts','Followup','Simulation','Source panel','Dictionary']
forecasts=[]
for vintage in ['harmonised','vintage']:
    boot={r['method']:r for r in read(f'results/{vintage}/paired_bootstrap.csv')}
    for r in read(f'results/{vintage}/test_metrics.csv'):
        forecasts.append(dict(dataset=vintage,**r)|boot[r['method']])
simulation=[r for r in read('results/simulation.csv') if float(r['phi'])==.85 and float(r['sigma2'])==6.25 and float(r['psi'])==0 and float(r['fraction'])==.2]
datasets={('Forecasts',7):forecasts,('Followup',7):read('results/harmonised/observational_followup.csv'),
    ('Followup',16):read('results/harmonised/rank_sensitivity.csv'),('Simulation',7):simulation,
    ('Source panel',7):read('data/processed/nhs_staff_survey_panel_2021_2025.csv')}
identifiers={'dataset','method','org_code','org_name','benchmark_group','legal_org_code','seed','error_kind'}
checked=0
for section in manifest['sheets']:
    key=(section['name'],section['first_row'])
    if key not in datasets:continue
    records=datasets[key];assert len(records)==section['data_rows']
    sheet=wb[section['name']]
    cells=list(sheet.iter_rows(min_row=section['first_row']+1,max_row=section['first_row']+len(records),max_col=len(section['columns'])))
    for row,record in zip(cells,records):
        for cell,field in zip(row,section['columns']):
            wanted=record[field];actual=cell.value
            assert cell.data_type!='f','Static snapshots must not contain formulas'
            if wanted=='':assert actual is None,(key,field,actual)
            elif field in identifiers:assert actual==wanted and isinstance(actual,str),(key,field,actual,wanted)
            elif wanted in {'True','False'}:assert actual==(wanted=='True')
            else:
                assert isinstance(actual,(float,int)) and not isinstance(actual,bool),(key,field,actual)
                assert math.isclose(actual,float(wanted),rel_tol=1e-12,abs_tol=1e-12),(key,field,actual,wanted)
            checked+=1
result=dict(status='passed',scientific_cells_compared=checked,sheets=wb.sheetnames,
    source_hashes_checked=len(manifest['sources']),seed_precision='exact strings preserved',
    missing_values='blanks preserved without zero substitution',static_values='numeric types verified; no formulas in scientific tables',
    file_sha256=sha(OUTPUT),validator_sha256=sha(Path(__file__)))
(ART/'saved_workbook_validation.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result,indent=2))
