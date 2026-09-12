"""Independently compare the saved workbook with all source CSVs and formulas.

By default require cached formula values (release validation). --formulas-only
allows an openpyxl rebuild before a spreadsheet engine has recalculated it.
"""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import math
from datetime import datetime, timezone
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def equivalent(actual, raw, key):
    if raw=='': return actual is None
    if raw in ['True','False']: return isinstance(actual,bool) and actual==(raw=='True')
    if key in {'state','org_code'}: return isinstance(actual,str) and actual==raw
    try:
        number=float(raw)
    except ValueError:
        return actual==raw
    return isinstance(actual,(int,float)) and not isinstance(actual,bool) and math.isclose(actual,number,rel_tol=1e-13,abs_tol=1e-14)


def main(formulas_only=False):
    path=ROOT/'results/nhs_aisi_results.xlsx'
    meta=json.loads((ROOT/'results/workbook_manifest.json').read_text())
    # Some engines omit optional worksheet dimension attributes. Loading the
    # modest saved tables normally obtains dimensions from actual cells.
    values=load_workbook(path,read_only=False,data_only=True)
    formulas=load_workbook(path,read_only=False,data_only=False)
    count=0
    for source in meta['source_files']:
        csvpath=ROOT/source['path']
        assert sha(csvpath)==source['sha256'],f'Source changed since workbook build: {source["path"]}'
        with csvpath.open(newline='') as f:
            reader=csv.DictReader(f);fields=reader.fieldnames;rows=list(reader)
        sheet=values[source['sheet']]
        assert sheet.max_row==len(rows)+1 and sheet.max_column==len(fields),(source['sheet'],'dimensions')
        assert [c.value for c in next(sheet.iter_rows())]==fields,(source['sheet'],'headers')
        for r,(cells,row) in enumerate(zip(sheet.iter_rows(min_row=2),rows),2):
            for key,cell in zip(fields,cells):
                assert equivalent(cell.value,row[key],key),(source['sheet'],r,key,cell.value,row[key])
                count+=1
    fcount=0
    for address,item in meta['formula_expectations'].items():
        sheet,cell=address.split('!')
        assert formulas[sheet][cell].value==item['formula'],(address,'formula')
        if not formulas_only:
            actual=values[sheet][cell].value; expected=item['value']
            assert (isinstance(actual,(int,float)) and math.isclose(actual,expected,rel_tol=1e-11,abs_tol=1e-12)) if isinstance(expected,(int,float)) else actual==expected,(address,actual,expected)
        fcount+=1
    errors=[]
    for sh in values:
        for row in sh:
            errors.extend(f'{sh.title}!{c.coordinate}:{c.value}' for c in row if c.data_type=='e')
    assert not errors,errors[:20]
    structure=load_workbook(path,read_only=False,data_only=False)
    assert not structure['Results'].sheet_view.showGridLines
    assert structure['Results'].freeze_panes=='C7'
    assert len(structure['Results'].conditional_formatting)==1
    assert len(structure['Exploratory'].merged_cells.ranges)==0
    assert set(values.sheetnames)==set(meta['sheets'])
    report=dict(status='passed',checked_at_utc=datetime.now(timezone.utc).isoformat(),
                source_cells_checked=count,formula_cells_checked=fcount,source_grids=len(meta['source_files']),
                sheets=len(values.sheetnames),cached_formula_values_checked=not formulas_only,
                native_excel_application_test=False,workbook_sha256=sha(path),source_hashes_unchanged=True,
                no_error_cells=True,verifier_sha256=sha(__file__))
    (ROOT/'results/workbook_validation.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--formulas-only',action='store_true')
    main(p.parse_args().formulas_only)
