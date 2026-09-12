"""Build a transparent Excel supplement from frozen result CSVs.

Run with Python and openpyxl. Excel recalculates linked comparisons on opening.
This script does not modify source results or run scientific experiments.
"""
from pathlib import Path
from datetime import datetime, timezone
import argparse
import csv
import hashlib
import json
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
SOURCES = [
    ("Synthetic", "synthetic.csv"), ("Stress", "synthetic_stress.csv"),
    ("Replay", "replay.csv"), ("Calibration", "replay_calibration.csv"),
    ("Allocations", "replay_allocations.csv"),
    ("NHS summary", "nhs_scenario_summary.csv"), ("NHS scenarios", "nhs_scenarios.csv"),
]
OPTIONAL = [("Exploratory", "replay_budget_extension.csv")]
TEXT_FIELDS = {"scenario", "policy", "state", "org_code", "org_name", "benchmark_group", "analysis_stage",
               "evidence_type", "preferred_channel"}
INT_FIELDS = {"budget", "audit_cost", "reps", "units", "k", "review_slots",
              "nominal_burnout_n", "action_priority_rank", "top20_validation_preferred"}
NAVY, LIGHT, INK = "183F59", "EEF3F6", "213747"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def typed(key, value):
    if value == "":
        return None
    if key in TEXT_FIELDS:
        return value
    if value in ("True", "False"):
        return value == "True"
    if key in INT_FIELDS:
        return int(value)
    try:
        val = float(value)
        if not math.isfinite(val):
            raise ValueError(f"Nonfinite source value: {key}={value}")
        return val
    except ValueError:
        return value


def style_table(sheet, first, columns, last, identifiers=1):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f"{get_column_letter(identifiers + 1)}{first + 1}"
    sheet.auto_filter.ref = f"A{first}:{get_column_letter(columns)}{last}"
    for row in sheet.iter_rows(min_row=first, max_row=last, max_col=columns):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(vertical="center", horizontal="right" if isinstance(cell.value, (int,float)) else "left")
            if isinstance(cell.value, (int,float)) and not isinstance(cell.value,bool):
                cell.number_format = "0.0000"
    for cell in sheet[first]:
        if cell.column > columns:
            break
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[first].height = 52
    sheet.print_title_rows = f"1:{first}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


DEFINITIONS = {
    "scenario": ("Named controlled sampling or data-generating mechanism; not an observed recruitment programme", "category"),
    "policy": ("Acquisition policy identifier; see policy definitions below", "category"),
    "bias_sd": ("Actual persistent unit-offset SD in synthetic world", "score units"),
    "assumed_beta": ("Persistent-offset SD assumed by inference, except explicitly bias-blind policies", "score units"),
    "audit_cost": ("Cost of a validation packet relative to an ordinary packet costing one", "packet-cost equivalents"),
    "budget": ("Maximum additional acquisition expenditure, excluding common initial sample", "packet-cost equivalents"),
    "reps": ("Paired independent Monte Carlo panels for that scenario; reused across policies/settings as documented", "panels"),
    "weighted": ("True: weighted finite-file outcome distribution and corresponding record-proportional-to-weight replay", "boolean"),
    "kappa": ("Injected ordinary-channel outcome tilt magnitude; sign is random by area and replicate", "dimensionless"),
    "validation_kappa": ("Injected second-channel tilt in the validation-channel stress", "dimensionless"),
    "units": ("Number of decision units: 40 synthetic, 48 BRFSS areas, or 190 NHS reporting units", "units"),
    "k": ("Number of selected BRFSS areas (9 of 48)", "areas"),
    "regret": ("Selected mean reference score minus oracle-selected mean reference score; lower is better", "score units"),
    "overlap": ("Fraction of selected units also in oracle selection; higher is better", "fraction"),
    "mse": ("Mean squared posterior-mean error across units", "squared score units"),
    "coverage": ("Fraction of unit targets inside nominal 95% model intervals; averaged across panels", "fraction"),
    "spent": ("Actual additional packet cost consumed", "packet-cost equivalents"),
    "validation_fraction": ("Validation expenditure divided by available budget", "fraction"),
    "regular_packets": ("Additional ordinary packets; each contains 20 sampled outcomes", "packets"),
    "validation_packets": ("Additional validation packets; each contains 20 sampled outcomes", "packets"),
    "posterior_sd": ("Average model posterior SD across units", "score units"),
    "repeat_floor": ("Mean model limiting target variance after arbitrarily many ordinary observations", "squared score units"),
    "regret_diff_uniform": ("Within-panel regret minus same-panel uniform-policy regret", "score units"),
    "regret_diff_ordinary_kg": ("Within-panel regret minus same-panel ordinary_kg regret", "score units"),
    "overlap_diff_uniform": ("Within-panel overlap minus same-panel uniform-policy overlap", "fraction"),
    "overlap_diff_ordinary_kg": ("Within-panel overlap minus same-panel ordinary_kg overlap", "fraction"),
    "state": ("Published state/DC code retained as text exactly as in the source CSV, including any zero-padding", "identifier"),
    "mu": ("2023 calibration: mean of area reference means", "score units"),
    "tau": ("2023 calibration: sample SD of area reference means", "score units"),
    "bias_mean": ("2023 mean offset over positive/negative synthetic recruitment tilts", "score units"),
    "ordinary_variance": ("2023 variance under ordinary sampling laws, averaged over unknown signs", "squared score units"),
    "validation_variance": ("2023 reference distribution variance assumed for second channel", "squared score units"),
    "ordinary_count_mean": ("Mean sampled outcome count in area, including the initial 20; draws can repeat", "outcome draws"),
    "validation_count_mean": ("Mean sampled outcome count in validation channel; draws can repeat", "outcome draws"),
    "mean_estimate": ("Mean final posterior target estimate across replay panels", "score units"),
    "reference": ("2024 complete-file area mean used only for evaluation", "score units"),
    "org_code": ("Official NHS organisation/reporting code", "identifier"),
    "org_name": ("Official NHS organisation name", "text"),
    "benchmark_group": ("NHS peer group defining within-group review slots", "category"),
    "ess_factor": ("Assumed multiplier on nominal response n; not an estimated survey effective sample size", "multiplier"),
    "assumed_tau": ("NHS scenario assumed between-unit target SD", "score units"),
    "assumed_individual_sd": ("NHS scenario assumed individual measurement SD", "score units"),
    "nominal_burnout_n": ("Published valid burnout response count; not survey effective sample size", "responses"),
    "observed_burnout": ("Published 2025 PP4_2 organisation score; higher means less burnout", "0-10 score"),
    "posterior_mean": ("Model-conditional NHS posterior target mean", "score units"),
    "repeat_floor_sd": ("Model limiting target SD under repeated ordinary measurements", "score units"),
    "selected": ("Model posterior-mean selection for fixed within-group review slots", "boolean"),
    "review_slots": ("Total fixed NHS review slots across peer groups", "slots"),
    "ordinary_KG": ("One-step expected improvement in selected mean score from ordinary channel, before dividing by cost", "score units"),
    "validation_KG": ("One-step expected improvement in selected mean score from second channel, before dividing by cost", "score units"),
    "ordinary_KG_per_cost": ("Ordinary one-step expected improvement divided by packet cost", "score / cost unit"),
    "validation_KG_per_cost": ("Validation one-step expected improvement divided by packet cost", "score / cost unit"),
    "preferred_channel": ("Larger model one-step improvement per cost; tied scores prefer ordinary", "category"),
    "action_priority_rank": ("Rank of model action value; not a validated risk or service-allocation rank", "rank"),
    "evidence_type": ("Explicit model-conditional status; no newly observed NHS outcomes", "text"),
    "median_nominal_n": ("Median nominal burnout valid-response count across 190 units", "responses"),
    "mean_posterior_sd": ("Mean model target SD across the 190 NHS units", "score units"),
    "mean_repeat_floor_sd": ("Mean model limiting target SD across NHS units", "score units"),
    "top20_validation_preferred": ("Count among top 20 action-priority units whose preferred channel is validation", "units"),
    "greatest_one_step_KG_per_cost": ("Largest model one-step improvement in selected mean score per cost over units/channels", "score / cost unit"),
    "analysis_stage": ("Post-primary exploratory budget extension, separately frozen after primary results were known", "category"),
}
POLICY = {
    "uniform": "Ordinary-only uniform allocation with bias-aware inference",
    "uniform_blind": "Ordinary-only uniform allocation with zero assumed persistent bias",
    "uncertainty": "Ordinary-channel uncertainty-based allocation with matched bias-aware inference",
    "boundary": "Ordinary-channel boundary/LUCB-inspired allocation; no inherited LUCB theorem",
    "bias_blind_kg": "Ordinary-channel KG with zero assumed persistent bias",
    "ordinary_kg": "Ordinary-channel cost-aware KG with matched bias-aware inference",
    "two_channel_uniform": "Uniform unit and equal-action-probability channel sampling",
    "balanced_mix": "Random channel mixture targeting equal expected cost shares before boundary effects",
    "validation_only": "Uniform-unit validation-only diagnostic, subject to affordability",
    "two_channel_kg": "One-step cost-normalised KG over both channels; not globally optimal",
}


def build(output):
    wb = Workbook(); wb.remove(wb.active)
    results = wb.create_sheet("Results"); results.sheet_properties.tabColor = NAVY
    src = {}; evidence = []
    for sheet_name, filename in SOURCES + [(s, f) for s,f in OPTIONAL if (ROOT/'results'/f).exists()]:
        path = ROOT/'results'/filename
        with path.open(newline="") as f:
            reader = csv.DictReader(f); fields = reader.fieldnames; raw = list(reader)
        typed_rows = [[typed(k, row[k]) for k in fields] for row in raw]
        sh = wb.create_sheet(sheet_name); sh.append(fields)
        for row in typed_rows: sh.append(row)
        style_table(sh, 1, len(fields), len(raw)+1, identifiers=1 if sheet_name in {"Calibration","NHS summary"} else 2)
        for j,key in enumerate(fields,1):
            sh.column_dimensions[get_column_letter(j)].width = 33 if key in {"scenario","policy"} else 23 if key in TEXT_FIELDS else 17
            if key in {"org_name", "evidence_type"}: sh.column_dimensions[get_column_letter(j)].width = 64
            if key == 'analysis_stage': sh.column_dimensions[get_column_letter(j)].width = 38
            if key == 'benchmark_group': sh.column_dimensions[get_column_letter(j)].width = 44
            if key in {'org_name','benchmark_group'}:
                for cell in sh[get_column_letter(j)][1:]:
                    cell.alignment = Alignment(vertical='center',wrap_text=True)
                    if len(str(cell.value)) > (62 if key == 'org_name' else 42):
                        sh.row_dimensions[cell.row].height = 32
            if key in INT_FIELDS:
                for col in sh.iter_cols(min_col=j,max_col=j,min_row=2):
                    for cell in col: cell.number_format = "#,##0"
            if key == 'state':
                for cell in sh[get_column_letter(j)][1:]: cell.number_format = '@'
            if 'KG' in key:
                for cell in sh[get_column_letter(j)][1:]: cell.number_format = '0.000E+00'
        src[sheet_name] = dict(fields=fields,raw=raw,typed=typed_rows,file=filename)
        evidence.append(dict(sheet=sheet_name,path=str(path.relative_to(ROOT)),sha256=sha(path),rows=len(raw),columns=len(fields)))

    results['A2'] = "Two-channel KG versus ordinary-channel KG"
    results['A2'].font = Font(name="Arial",size=14,bold=True,color=NAVY)
    results['A3'] = "All available paired settings. Positive regret difference is worse. Intervals describe Monte Carlo uncertainty, not population uncertainty."
    results['A3'].font = Font(name="Arial",size=10,italic=True,color=INK)
    headers = ["Study","Scenario","Bias SD","Validation cost","Budget","Replicates","Ordinary KG regret","Two-channel KG regret","Regret difference","Paired MCSE","95% MC lower","95% MC upper","Relative regret change","Validation cost share","Direction from MC interval"]
    for c,h in enumerate(headers,1): results.cell(6,c,h)
    expected = {}; r = 7
    for sheet_name in ['Synthetic','Stress','Replay','Exploratory']:
        if sheet_name not in src: continue
        source = src[sheet_name]; fields = source['fields']; rows=source['raw']
        config = [k for k in fields if k not in {'policy','reps'} and not any(k.endswith('_'+s) for s in ['mean','mcse','lo','hi'])]
        bykey = {}
        for j,row in enumerate(rows,2):
            key=tuple(row[k] for k in config)+(row['policy'],)
            if key in bykey: raise ValueError(f'Duplicate configuration: {key}')
            bykey[key]=j
        for j,row in enumerate(rows,2):
            if row['policy']!='two_channel_kg': continue
            baseline=bykey[tuple(row[k] for k in config)+('ordinary_kg',)]
            def ref(key, rr=j):
                return f"'{sheet_name}'!{get_column_letter(fields.index(key)+1)}{rr}"
            for col,value in enumerate([sheet_name,row['scenario'],typed('bias_sd',row.get('bias_sd','')),
                                       int(row['audit_cost']),int(row['budget']),int(row['reps'])],1):results.cell(r,col,value)
            formulas = {7:f'={ref("regret_mean",baseline)}',8:f'={ref("regret_mean")}',9:f'=H{r}-G{r}',
                        10:f'={ref("regret_diff_ordinary_kg_mcse")}',11:f'=I{r}-1.96*J{r}',12:f'=I{r}+1.96*J{r}',
                        13:f'=IF(G{r}=0,"n.a.",I{r}/G{r})',14:f'={ref("validation_fraction_mean")}',
                        15:f'=IF(K{r}>0,"Higher regret",IF(L{r}<0,"Lower regret","Interval crosses zero"))'}
            base=float(rows[baseline-2]['regret_mean']);new=float(row['regret_mean']);diff=new-base;se=float(row['regret_diff_ordinary_kg_mcse'])
            assert math.isclose(diff,float(row['regret_diff_ordinary_kg_mean']),abs_tol=1e-12)
            vals={7:base,8:new,9:diff,10:se,11:diff-1.96*se,12:diff+1.96*se,13:'n.a.' if base==0 else diff/base,
                  14:float(row['validation_fraction_mean']),15:'Higher regret' if diff-1.96*se>0 else 'Lower regret' if diff+1.96*se<0 else 'Interval crosses zero'}
            for col,formula in formulas.items():
                addr=f'{get_column_letter(col)}{r}';results[addr]=formula
                expected[f'Results!{addr}']=dict(formula=formula,value=vals[col])
            r+=1
    style_table(results,6,15,r-1,identifiers=2)
    widths=[18,29,12,14,12,12,16,17,16,14,16,16,17,16,26]
    for c,w in enumerate(widths,1):results.column_dimensions[get_column_letter(c)].width=w
    for rr in range(7,r):
        for c in range(4,7):results.cell(rr,c).number_format='#,##0'
        for c in [13,14]:results.cell(rr,c).number_format='0.0%'
    for rr in range(7,r):
        for c in range(7,15):results.cell(rr,c).alignment=Alignment(horizontal='right',vertical='center')
    results.conditional_formatting.add(f'I7:I{r-1}',CellIsRule(operator='greaterThan',formula=['0'],fill=PatternFill('solid',fgColor='FBE9E7')))
    results.conditional_formatting.add(f'I7:I{r-1}',CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='E6EFF8')))

    read=wb.create_sheet('ReadMe');read.sheet_properties.tabColor='7D8C95'
    read.append(['Topic / variable','Definition','Units / interpretation','Source'])
    notes=[
        ('Purpose','Frozen research-result supplement. Source CSVs and executed scripts are the scientific authority.','No NHS trial or deployment','README.md'),
        ('Data','420,249 retained BRFSS records: 206,098 in 2023 and 214,151 in 2024, across 47 common states plus DC.','Independent cross-sections; 48 areas','docs/microdata_audit.md'),
        ('Outcome','BRFSS score=(30-mentally-unhealthy-days)/3. NHS PP4_2 is a different composite.','0-10 higher is better','docs/analysis_protocol.md'),
        ('Synthetic replication','2,000 panels per world; 4 base worlds and 7 stress worlds. Base worlds reused across 24 cost-budget settings.','Do not count shared panels as independent','results/synthetic_manifest.json'),
        ('Replay replication','1,000 panels per primary scenario, 5 scenarios and 18 cost-budget settings. Each policy uses paired outcome streams.','With-replacement finite-file sampling','results/replay_manifest.json'),
        ('Formula intervals','Results uses mean difference and its paired MCSE; interval=mean +/- 1.96*MCSE. It never combines independent-policy SEs.','Approximate 95% Monte Carlo interval','experiments/common.py'),
        ('Relative change','(two-channel regret - ordinary regret)/ordinary regret; unavailable for a zero baseline.','Descriptive ratio; no ratio CI is claimed','Results formulas'),
        ('NHS scenario status','190 reporting units, 189 identified legal entities, 36 peer-group review slots. Beta and precision multipliers are assumptions.','No new NHS measurement or estimated survey ESS','results/nhs_scenario_summary.csv'),
        ('New results','Rerun this builder after result sources change. Formula values recalculate on opening in Excel. Published copy was also recalculated and checked.','No editable research model is embedded','artifact_tools/build_results.py'),
        ('Privacy and rights','Public agency data retain source rights. No endorsement by NHS, CDC, HHS or the United States Government is implied.','Do not identify individuals','docs/microdata_audit.md'),
        ('Suffix: _mean','Average metric across paired independent panels for that setting.','Metric units','experiments/common.py'),
        ('Suffix: _mcse','Sample standard deviation over independent panels divided by sqrt(reps).','Metric units','experiments/common.py'),
        ('Suffix: _lo / _hi','Mean minus/plus 1.96 MCSE. Coverage-rate intervals are MC intervals on average coverage.','Metric units','experiments/common.py'),
        ('Calibration bias_sd','In Calibration only: 2023 SD of synthetic positive/negative tilt offsets, not the synthetic-world input.','Score units','experiments/replay.py'),
        ('Secondary weighted replay','Draws proportional to public survey weight target the weighted complete-file mean. Not a subset ratio or new population-confidence analysis.','Controlled finite-file sensitivity','experiments/replay.py'),
    ]
    if 'Exploratory' in src:
        notes.append(('Budget extension','Post-primary exploratory larger-budget replay; distinct from frozen primary 1,000-replicate study.','500 panels; see extension protocol','results/replay_budget_extension.csv'))
    for row in notes:read.append(row)
    for key,(definition,unit) in DEFINITIONS.items():read.append([key,definition,unit,'Source CSV field / executed script'])
    for key,definition in POLICY.items():read.append([key,definition,'Policy identifier','experiments/allocation.py'])
    for item in evidence:read.append([item['sheet'],f"{item['rows']} rows x {item['columns']} columns; SHA256 {item['sha256']}",'Frozen source grid',item['path']])
    read.append(['BRFSS 2023','CDC annual public-use data; original source manifest retains URLs and hashes.','Source data','https://www.cdc.gov/brfss/annual_data/annual_2023.html'])
    read.append(['BRFSS 2024','CDC annual public-use data; original source manifest retains URLs and hashes.','Source data','https://www.cdc.gov/brfss/annual_data/annual_2024.html'])
    style_table(read,1,4,read.max_row)
    for c,w in zip('ABCD',[30,105,42,64]):read.column_dimensions[c].width=w
    for row in read.iter_rows(min_row=2):
        read.row_dimensions[row[0].row].height=42
        for cell in row:cell.alignment=Alignment(vertical='center',wrap_text=True)
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
    wb.properties.creator='Shengwei Zhang';wb.properties.title='Which Uncertainty Is Worth Reducing? Budgeted Measurement for Wellbeing Prioritisation'
    output.parent.mkdir(parents=True,exist_ok=True);wb.save(output)
    meta=dict(created_at_utc=datetime.now(timezone.utc).isoformat(),source_files=evidence,
              sheets={s.title:{'rows':s.max_row,'columns':s.max_column} for s in wb},
              authoring='openpyxl, explicitly requested; artifact-tool used for recalculation and visual QA',
              builder_sha256=sha(__file__),formula_expectations=expected)
    (ROOT/'results/workbook_manifest.json').write_text(json.dumps(meta,indent=2)+'\n')
    print(f'Created {output}; {len(wb.sheetnames)} sheets; {len(expected)} transparent formula cells')


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,default=ROOT/'results/nhs_aisi_results.xlsx')
    build(parser.parse_args().output)
